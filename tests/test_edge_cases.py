"""
엣지케이스 회귀 테스트. 스킬을 고친 뒤 반드시 돌린다.

    python tests/test_edge_cases.py

여기 있는 케이스는 전부 '조용히 틀린 값이 나올 수 있었던' 실제 상황이다.
크래시보다 그럴듯한 오답이 더 위험하므로, 값까지 확인한다.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import date

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(SKILL, "scripts"))
FIX = os.path.join(HERE, "fixtures")

from common import (load_config, load_uploads, match_offsets, live_rows,  # noqa: E402
                    month_range, is_in_grace, biz_checksum_ok, norm_biz,
                    deadline_for)

fails = []


def check(name, cond, detail=""):
    print(f"  {'ok  ' if cond else 'FAIL'}  {name}" + (f"  — {detail}" if detail else ""))
    if not cond:
        fails.append(name)


def run(args, cwd=SKILL):
    return subprocess.run([sys.executable] + args, cwd=cwd,
                          capture_output=True, text=True)


class sandbox:
    """
    저장소 트리를 통째로 복사한 임시 작업본.

    이전 테스트는 실제 config/vendors.json 과 state/last-run.json 을 덮어썼다가
    finally 에서 되돌렸다. 중간에 죽으면 vendors.json 이 빈 목록으로 남고,
    push 의 기본 묶음에 vendors 가 들어 있어서 그대로 저장소에 올라간다.
    state 에는 미래 날짜가 남아 이후 push 가 역행 가드에 전부 막힌다.
    복사본에서 돌리면 실제 파일을 건드릴 일이 없다.
    """

    def __enter__(self):
        self.dir = tempfile.mkdtemp(prefix="hometax-test-")
        for name in ["config", "references", "scripts", "state", "tests"]:
            src = os.path.join(SKILL, name)
            if os.path.isdir(src):
                shutil.copytree(src, os.path.join(self.dir, name))
        for name in ["sync.py", "README.md", "SKILL.md"]:
            src = os.path.join(SKILL, name)
            if os.path.exists(src):
                shutil.copy(src, self.dir)
        return self.dir

    def __exit__(self, *a):
        shutil.rmtree(self.dir, ignore_errors=True)


# ---------------------------------------------------------------- 단위

def t_biz():
    print("\n[사업자번호]")
    check("체크섬 통과", biz_checksum_ok("239-86-03249"))
    check("체크섬 오타 탐지", not biz_checksum_ok("239-86-03248"))
    check("전각 숫자 정규화", norm_biz("２３９８６０３２４９") == "2398603249")
    check("자릿수 부족 거부", not biz_checksum_ok("12345"))


def t_months():
    print("\n[월 범위 · 유예]")
    check("1월 실행시 전년 12월 + 1월",
          month_range(date(2027, 1, 5)) == [(2026, 12), (2027, 1)])
    check("10월 실행시 10개월", len(month_range(date(2026, 10, 2))) == 10)
    check("12월 실행시 12개월", len(month_range(date(2026, 12, 31))) == 12)
    check("2월 실행에는 전년이 안 붙음",
          month_range(date(2026, 2, 3)) == [(2026, 1), (2026, 2)])
    cfg = load_config()
    check("12월분 기한은 익년 1월", deadline_for(2026, 12, 10) == date(2027, 1, 10))
    check("9/6 기준 8월분은 유예", is_in_grace(2026, 8, date(2026, 9, 6), cfg))
    check("9/15 기준 8월분은 확정", not is_in_grace(2026, 8, date(2026, 9, 15), cfg))
    cfg2 = dict(cfg)
    cfg2["grace"] = dict(cfg["grace"], mode="strict")
    check("strict 모드는 유예 없음", not is_in_grace(2026, 8, date(2026, 9, 6), cfg2))


def t_empty_frames():
    print("\n[빈 데이터 · 전액취소]")
    cfg = load_config()
    df, _ = load_uploads(FIX, cfg)
    match_offsets(df, cfg)
    check("빈 DataFrame 에 live_rows", len(live_rows(df.iloc[0:0])) == 0)

    # 전 건이 상계된 거래처: 최종수취일이 취소일로 새어나오면 안 된다
    g = df[(df["공급자번호"] == "2120459010") & (df["월"] >= 4)]
    lv = live_rows(g)
    check("전액취소 거래처는 유효건 0", len(lv) == 0)
    last = lv["작성일"].max().date() if len(lv) else None
    check("전액취소 거래처 최종수취일은 None", last is None,
          "취소된 날짜를 최종수취일로 쓰면 조용한 오답")

    only_neg = df[df["공급가액n"] < 0].head(3)
    check("마이너스만 있는 집합도 처리", len(live_rows(only_neg)) >= 0)


def t_kind_detection():
    print("\n[파일 종류 판정]")
    cfg = load_config()
    with tempfile.TemporaryDirectory() as d:
        # 파일명을 완전히 바꿔도 컬럼 구성으로 맞춰야 한다
        for i, f in enumerate(sorted(os.listdir(FIX))):
            shutil.copy(os.path.join(FIX, f), os.path.join(d, f"zzz_{i}.xls"))
        df, meta = load_uploads(d, cfg)
        kinds = sorted({m["kind"] for m in meta})
        check("파일명 무관 세/계 판정", kinds == ["계", "세"],
              f"판정 결과 {kinds}")


def t_offset_no_miss_match():
    print("\n[상계 오매칭]")
    cfg = load_config()
    df, _ = load_uploads(FIX, cfg)
    pairs = match_offsets(df, cfg)
    check("상계쌍이 잡힘", pairs > 0, f"{pairs}쌍")
    paired = df[df["상계"].str.contains("상계")]
    # 매월 동일 금액인 케이티가 잘못 엮이지 않았는지
    kt = paired[paired["공급자번호"] == "1028142945"]
    check("동일금액 반복 거래처 오매칭 없음", len(kt) == 0)
    for _, r in paired[paired["상계"] == "취소분(상계)"].iterrows():
        mate = paired[(paired["공급자번호"] == r["공급자번호"]) &
                      (paired["공급가액n"] == -r["공급가액n"]) &
                      (paired["상계"] == "원본(상계됨)")]
        check(f"상계쌍 품목 일치 {r['상호']}",
              len(mate) and str(mate.iloc[0].get("품목명")) == str(r.get("품목명")))
        break


# ---------------------------------------------------------------- 통합

def t_run_variants():
    print("\n[실행 변형]")
    with tempfile.TemporaryDirectory() as out:
        base = ["scripts/build_report.py", "--uploads", FIX, "--out", out, "--no-state"]

        r = run(base + ["--as-of", "2026-09-06"])
        check("기본 실행", r.returncode == 0, r.stderr.strip()[-200:])

        r2 = run(base + ["--as-of", "2026-09-06", "--strict"])
        check("strict 모드 실행", r2.returncode == 0, r2.stderr.strip()[-200:])
        check("strict 는 '기한 전' 0건", "기한 전 0" in r2.stdout, r2.stdout.strip()[-120:])

        r3 = run(base + ["--as-of", "2026-01-15"])
        check("1월 초 실행(대상 1개월)", r3.returncode == 0, r3.stderr.strip()[-200:])
        check("1월 실행시 전년 12월 포함",
              "전년 12월~1월" in r3.stdout, r3.stdout.strip()[:80])

        r4 = run(base + ["--as-of", "2026-12-31"])
        check("연말 실행(대상 12개월)", r4.returncode == 0 and "1월~12월" in r4.stdout)


def t_single_kind():
    print("\n[세금계산서만 / 계산서만]")
    cfg = load_config()
    for keep, label in [("세금", "세금계산서만"), ("계산서목록", "계산서만")]:
        with tempfile.TemporaryDirectory() as d, tempfile.TemporaryDirectory() as out:
            for f in os.listdir(FIX):
                if (keep in f) if keep == "세금" else ("세금" not in f):
                    shutil.copy(os.path.join(FIX, f), os.path.join(d, f))
            n = len(os.listdir(d))
            r = run(["scripts/build_report.py", "--uploads", d, "--out", out,
                     "--no-state", "--as-of", "2026-09-06"])
            check(f"{label} ({n}개) 실행", r.returncode == 0, r.stderr.strip()[-200:])


def t_empty_vendors():
    print("\n[vendors.json 이 빈 경우]")
    with sandbox() as sb, tempfile.TemporaryDirectory() as out:
        json.dump({"vendors": []},
                  open(os.path.join(sb, "config", "vendors.json"), "w",
                       encoding="utf-8"))
        r = run(["scripts/build_report.py", "--uploads", FIX, "--out", out,
                 "--no-state", "--as-of", "2026-09-06"], cwd=sb)
        check("빈 목록으로도 실행됨", r.returncode == 0, r.stderr.strip()[-200:])
        check("전부 대상외로 분류", "점검대상 0곳" in r.stdout, r.stdout.strip()[-160:])
        r2 = run(["scripts/check_input.py", FIX, "--as-of", "2026-09-06"], cwd=sb)
        check("check_input 이 빈 목록 경고", "비어 있음" in r2.stdout)


def t_broken_json():
    """설정·이력 파일이 깨졌을 때 트레이스백 대신 안내가 나오는지."""
    print("\n[JSON 손상]")
    cases = [("config/vendors.json", "거래처 목록"),
             ("config/check-config.json", "설정"),
             ("state/last-run.json", "이력")]
    for rel, word in cases:
        with sandbox() as sb, tempfile.TemporaryDirectory() as out:
            open(os.path.join(sb, rel), "w", encoding="utf-8").write("{ 깨짐")
            r = run(["scripts/build_report.py", "--uploads", FIX, "--out", out,
                     "--as-of", "2026-09-06"], cwd=sb)
            blob = r.stdout + r.stderr
            check(f"{rel} 손상시 트레이스백 없음",
                  "Traceback" not in blob, blob.strip().splitlines()[-1:] or [""])
            check(f"{rel} 손상시 안내 메시지", "[중단]" in blob)


def t_missing_vendor_in_data():
    print("\n[자료에 한 건도 없는 거래처]")
    with sandbox() as sb, tempfile.TemporaryDirectory() as out:
        p = os.path.join(sb, "config", "vendors.json")
        d = json.load(open(p, encoding="utf-8"))
        d["vendors"].append({"biz_no": "1048119147", "name": "테스트 유령업체",
                             "cycle": "매월", "note": "엣지케이스 테스트용"})
        json.dump(d, open(p, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
        r = run(["scripts/build_report.py", "--uploads", FIX, "--out", out,
                 "--no-state", "--as-of", "2026-09-06"], cwd=sb)
        check("0건 거래처 있어도 실행", r.returncode == 0, r.stderr.strip()[-300:])
        files = os.listdir(out)
        if files:
            from openpyxl import load_workbook
            wb = load_workbook(os.path.join(out, files[0]))
            ws = wb[load_config()["sheets"]["missing"]]
            txt = "\n".join(str(c.value) for row in ws.iter_rows() for c in row)
            check("0건 거래처가 '한 건도 없음'으로 잡힘", "한 건도 없음" in txt)
            check("유령업체가 목록에 등장", "테스트 유령업체" in txt)


def t_state_cycle():
    print("\n[직전 실행 비교]")
    with sandbox() as sb, tempfile.TemporaryDirectory() as out:
        sp = os.path.join(sb, "state", "last-run.json")
        if os.path.exists(sp):
            os.remove(sp)
        r1 = run(["scripts/build_report.py", "--uploads", FIX, "--out", out,
                  "--as-of", "2026-09-06"], cwd=sb)
        check("1회차 실행", r1.returncode == 0, r1.stderr.strip()[-200:])
        check("상태 파일 생성", os.path.exists(sp))
        first = json.load(open(sp, encoding="utf-8"))
        check("flagged 기록됨", len(first["flagged"]) > 0, f"{len(first['flagged'])}건")
        check("first_seen 기록됨", all("first_seen" in x for x in first["flagged"]))

        r2 = run(["scripts/build_report.py", "--uploads", FIX, "--out", out,
                  "--as-of", "2026-10-06"], cwd=sb)
        check("2회차 실행", r2.returncode == 0, r2.stderr.strip()[-200:])
        from openpyxl import load_workbook
        f = sorted(os.listdir(out))[-1]
        ws = load_workbook(os.path.join(out, f))[load_config()["sheets"]["missing"]]
        txt = "\n".join(str(c.value) for row in ws.iter_rows() for c in row)
        check("'계속' 상태 표시됨", "계속" in txt)


def _clipped_facts(as_of, vendors=None, strict=False):
    """업로드가 실행일까지만 있는 현실 상황을 재현해 등급을 매긴다."""
    import copy
    import build_report as B
    from common import load_vendors
    cfg = copy.deepcopy(load_config())
    if strict:
        cfg["grace"]["mode"] = "strict"
    df, _ = load_uploads(FIX, cfg)
    match_offsets(df, cfg)
    df = df[df["받는자번호"] == cfg["my_biz"]["biz_no"]]
    df = df[df["작성일"].dt.date <= as_of]
    facts, _ = B.build_facts(df, vendors or load_vendors(), as_of, cfg)
    return B.grade_all(facts, as_of, cfg)


def _n_action(rows):
    return sum(1 for r in rows if r["grade"] == "확인 필요")


def t_month_boundary():
    """실행 시점이 월초일 때 전 거래처가 오탐으로 잡히지 않는지."""
    print("\n[월초·연초 오탐]")
    from common import load_vendors
    total = len(load_vendors())

    n_jan = _n_action(_clipped_facts(date(2026, 1, 5)))
    check("1월 초 실행이 전 거래처를 확인 필요로 만들지 않음", n_jan < total // 2,
          f"{total}곳 중 확인 필요 {n_jan}건")

    n_feb = _n_action(_clipped_facts(date(2026, 2, 5)))
    check("2월 초 실행도 오탐 없음", n_feb == 0, f"확인 필요 {n_feb}건")

    # 정상 동작하던 시점은 그대로여야 한다
    n_sep = _n_action(_clipped_facts(date(2026, 9, 6)))
    check("9/6 기준 판정은 유지", 0 < n_sep < total // 2, f"확인 필요 {n_sep}건")


def t_strict_current_month():
    """strict 가 아직 끝나지 않은 당월을 미수취로 세면 안 된다."""
    print("\n[strict 당월 처리]")
    cfg = load_config()
    cfg2 = dict(cfg)
    cfg2["grace"] = dict(cfg["grace"], mode="strict")
    check("strict 라도 실행 당월은 유예",
          is_in_grace(2026, 9, date(2026, 9, 6), cfg2))
    check("strict 는 지난 달을 유예하지 않음",
          not is_in_grace(2026, 8, date(2026, 9, 6), cfg2))

    # strict 는 '기한 전'(전월분 대기)을 확인 필요로 승격시키는 것까지가 정상이다.
    # 그보다 더 늘면 아직 끝나지도 않은 당월을 결번으로 세고 있다는 뜻이다.
    rows_auto = _clipped_facts(date(2026, 9, 6))
    auto = _n_action(rows_auto)
    grace = sum(1 for r in rows_auto if r["grade"] == "기한 전")
    strict = _n_action(_clipped_facts(date(2026, 9, 6), strict=True))
    check("strict 가 당월 결번으로 전 거래처를 잡지 않음", strict <= auto + grace,
          f"auto {auto}건(+기한 전 {grace}) → strict {strict}건")


def t_until_last_year():
    """작년에 끝난 거래처(until=작년)는 올해 점검 대상이 아니다."""
    print("\n[until 이 작년인 거래처]")
    import build_report as B
    from common import load_vendors
    check("_m: until 이 작년이면 상한 0", B._m("2025-08", 2026, 12) == 0,
          f"실제 {B._m('2025-08', 2026, 12)}")
    check("_m: since 가 작년이면 하한 1", B._m("2025-08", 2026, 1) == 1)
    check("_m: since 가 내년이면 하한 13", B._m("2027-03", 2026, 1) == 13)
    check("_m: until 이 내년이면 상한 12", B._m("2027-03", 2026, 12) == 12)

    v = load_vendors() + [{"biz_no": "1048119147", "name": "작년종료테스트",
                           "cycle": "매월", "note": "", "until": "2025-08"}]
    rows = _clipped_facts(date(2026, 9, 6), vendors=v)
    hit = [r for r in rows if r["name"] == "작년종료테스트"]
    check("작년에 끝난 거래처는 확인 필요로 안 잡힘",
          not any(r["grade"] == "확인 필요" for r in hit),
          f"{[r['grade'] for r in hit] or '목록에 없음'}")


def t_sheet_headers():
    """헤더 행 서식과 틀고정이 실제 헤더 행에 맞는지."""
    print("\n[헤더 행·틀고정]")
    from openpyxl import load_workbook
    cfg = load_config()
    with tempfile.TemporaryDirectory() as out:
        r = run(["scripts/build_report.py", "--uploads", FIX, "--out", out,
                 "--no-state", "--as-of", "2026-09-06"])
        check("실행", r.returncode == 0, r.stderr.strip()[-200:])
        wb = load_workbook(os.path.join(out, os.listdir(out)[0]))
        for key, must in [("matrix", "공급자번호"), ("missing", "등급"),
                          ("unlisted", "추가 권장"), ("raw", "승인번호")]:
            ws = wb[cfg["sheets"][key]]
            hdr = next((i for i in range(1, 9)
                        if any(str(c.value).strip() == must for c in ws[i])), None)
            check(f"{cfg['sheets'][key]}: 헤더 행을 찾음", hdr is not None)
            if hdr is None:
                continue
            filled = all(c.fill and c.fill.fill_type for c in ws[hdr]
                         if c.value is not None)
            check(f"{cfg['sheets'][key]}: 헤더 행에 서식이 칠해짐", filled)
            frozen = int("".join(ch for ch in (ws.freeze_panes or "") if ch.isdigit()))
            check(f"{cfg['sheets'][key]}: 틀고정이 헤더 아래", frozen == hdr + 1,
                  f"헤더 {hdr}행 / 틀고정 {ws.freeze_panes}")


def t_duplicate_uploads():
    """같은 구간 파일을 두 번 올린 산출물을 validate 가 잡는지."""
    print("\n[중복 업로드]")
    with tempfile.TemporaryDirectory() as d, tempfile.TemporaryDirectory() as out:
        names = sorted(os.listdir(FIX))
        for f in names:
            shutil.copy(os.path.join(FIX, f), os.path.join(d, f))
        shutil.copy(os.path.join(FIX, names[-1]), os.path.join(d, "중복본.xls"))
        r = run(["scripts/build_report.py", "--uploads", d, "--out", out,
                 "--no-state", "--as-of", "2026-09-06"])
        check("중복 입력으로도 리포트는 생성됨", r.returncode == 0)
        x = os.path.join(out, os.listdir(out)[0])
        v = run(["scripts/validate.py", x, "--uploads", d, "--as-of", "2026-09-06"])
        check("validate 가 중복을 FAIL 로 잡음",
              v.returncode == 1 and "승인번호 중복" in v.stdout,
              v.stdout.strip().splitlines()[-1] if v.stdout else "")


def t_validate_strict():
    """strict 산출물을 --strict 로 검증하면 통과해야 한다."""
    print("\n[validate --strict]")
    with tempfile.TemporaryDirectory() as out:
        r = run(["scripts/build_report.py", "--uploads", FIX, "--out", out,
                 "--no-state", "--as-of", "2026-09-06", "--strict"])
        check("strict 리포트 생성", r.returncode == 0, r.stderr.strip()[-200:])
        x = os.path.join(out, os.listdir(out)[0])
        v = run(["scripts/validate.py", x, "--uploads", FIX,
                 "--as-of", "2026-09-06", "--strict"])
        check("--strict 로 검증하면 FAIL 없음", v.returncode == 0,
              v.stdout.strip().splitlines()[-1] if v.stdout else v.stderr[-200:])


def t_december():
    """전년 12월분이 점검 기간에 들어오는지 (연초 실행의 핵심)."""
    print("\n[전년 12월]")
    with tempfile.TemporaryDirectory() as out:
        r = run(["scripts/build_report.py", "--uploads", FIX, "--out", out,
                 "--no-state", "--as-of", "2026-01-05"])
        check("1월 초 실행", r.returncode == 0, r.stderr.strip()[-200:])
        from openpyxl import load_workbook
        cfg = load_config()
        ws = load_workbook(os.path.join(out, os.listdir(out)[0]))[cfg["sheets"]["matrix"]]
        hdr = next(i for i in range(1, 9)
                   if any(str(c.value).strip() == "공급자번호" for c in ws[i]))
        cols = [str(c.value) for c in ws[hdr] if c.value]
        check("매트릭스에 '전년 12월' 열이 있음", "전년 12월" in cols, str(cols[:6]))
        check("D-day 안내에 전년 12월 기한", "전년 12월분 기한" in r.stdout,
              r.stdout.strip()[-160:])


def t_as_of_clipping():
    """--as-of 가 그 시점 이후 데이터를 실제로 잘라내는지."""
    print("\n[--as-of 절단]")
    with tempfile.TemporaryDirectory() as out:
        r = run(["scripts/build_report.py", "--uploads", FIX, "--out", out,
                 "--no-state", "--as-of", "2026-03-05"])
        check("과거 시점 실행", r.returncode == 0, r.stderr.strip()[-200:])
        check("기간 밖 건수가 보고됨", "기간 밖 제외" in r.stdout)
        from openpyxl import load_workbook
        cfg = load_config()
        wb = load_workbook(os.path.join(out, os.listdir(out)[0]))
        ws = wb[cfg["sheets"]["raw"]]
        hdr = {str(c.value).strip(): c.column for c in ws[1] if c.value}
        late = [str(ws.cell(i, hdr["작성일자"]).value)
                for i in range(2, ws.max_row + 1)
                if str(ws.cell(i, hdr["작성일자"]).value) > "2026-03-05"]
        check("원본정제데이터에 기준일 이후 행이 없음", not late, str(late[:3]))
        mx = wb[cfg["sheets"]["matrix"]]
        h = next(i for i in range(1, 9)
                 if any(str(c.value).strip() == "공급자번호" for c in mx[i]))
        cols = [str(c.value) for c in mx[h] if c.value]
        check("4월 이후 열이 없음", "4월" not in cols, str(cols[:8]))


def t_validate_catches_regressions():
    """검증기가 실제로 뭔가를 잡는지. 산출물을 일부러 훼손해 확인한다."""
    print("\n[검증기가 실제로 잡는가]")
    from openpyxl import load_workbook
    cfg = load_config()
    with tempfile.TemporaryDirectory() as out:
        run(["scripts/build_report.py", "--uploads", FIX, "--out", out,
             "--no-state", "--as-of", "2026-09-06"])
        good = os.path.join(out, os.listdir(out)[0])
        base = run(["scripts/validate.py", good, "--uploads", FIX,
                    "--as-of", "2026-09-06"])
        check("정상 산출물은 통과", base.returncode == 0,
              base.stdout.strip().splitlines()[-1] if base.stdout else "")

        # 매트릭스 한 칸을 바꾸면 월별 합 대사가 깨져야 한다
        broken = os.path.join(out, "broken.xlsx")
        wb = load_workbook(good)
        mx = wb[cfg["sheets"]["matrix"]]
        h = next(i for i in range(1, 9)
                 if any(str(c.value).strip() == "공급자번호" for c in mx[i]))
        for i in range(h + 1, mx.max_row + 1):
            for j in range(4, mx.max_column):
                if isinstance(mx.cell(i, j).value, int) and mx.cell(i, j).value > 0:
                    mx.cell(i, j).value += 1000
                    break
            else:
                continue
            break
        wb.save(broken)
        r = run(["scripts/validate.py", broken, "--uploads", FIX,
                 "--as-of", "2026-09-06"])
        check("금액을 바꾸면 월별 합 대사가 FAIL",
              r.returncode == 1 and "월별 합 대사" in r.stdout)

        # 대상외 시트를 지우면 그 시트 검사가 FAIL 이어야 한다
        broken2 = os.path.join(out, "broken2.xlsx")
        wb2 = load_workbook(good)
        u = wb2[cfg["sheets"]["unlisted"]]
        u.delete_rows(4, u.max_row)
        wb2.save(broken2)
        r2 = run(["scripts/validate.py", broken2, "--uploads", FIX,
                  "--as-of", "2026-09-06"])
        check("대상외 시트를 비우면 FAIL",
              r2.returncode == 1 and "대상외공급자 시트" in r2.stdout)


def t_config_single_source():
    """단일 출처 검사가 진짜로 하드코딩을 잡는지 (예전엔 영영 0건이었다)."""
    print("\n[단일 출처 검사]")
    with sandbox() as sb, tempfile.TemporaryDirectory() as out:
        run(["scripts/build_report.py", "--uploads", FIX, "--out", out,
             "--no-state", "--as-of", "2026-09-06"], cwd=sb)
        x = os.path.join(out, os.listdir(out)[0])
        ok = run(["scripts/validate.py", x, "--uploads", FIX,
                  "--as-of", "2026-09-06"], cwd=sb)
        check("깨끗한 상태에서는 통과", ok.returncode == 0,
              ok.stdout.strip().splitlines()[-1] if ok.stdout else "")

        # 월 범위를 고정값으로 박으면 잡아야 한다
        cp = os.path.join(sb, "scripts", "common.py")
        src = open(cp, encoding="utf-8").read()
        open(cp, "w", encoding="utf-8").write(
            src.replace("months = [(as_of.year, m) for m in range(1, as_of.month + 1)]",
                        "months = [(as_of.year, m) for m in range(1, 10)]"))
        r = run(["scripts/validate.py", x, "--uploads", FIX,
                 "--as-of", "2026-09-06"], cwd=sb)
        check("월 범위 고정값을 잡음",
              "월 범위 고정값" in r.stdout, r.stdout.strip()[-120:])

        # config 에 아무도 안 읽는 키를 넣으면 잡아야 한다
        shutil.copy(cp + ".bak", cp) if os.path.exists(cp + ".bak") else None
        open(cp, "w", encoding="utf-8").write(src)
        cfgp = os.path.join(sb, "config", "check-config.json")
        d = json.load(open(cfgp, encoding="utf-8"))
        d["thresholds"]["nobody_reads_this"] = 42
        json.dump(d, open(cfgp, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
        r2 = run(["scripts/validate.py", x, "--uploads", FIX,
                  "--as-of", "2026-09-06"], cwd=sb)
        check("죽은 config 키를 잡음",
              "죽은 키" in r2.stdout and "nobody_reads_this" in r2.stdout,
              r2.stdout.strip()[-120:])


def t_cycle_warning():
    """cycle 이 실제 패턴과 어긋나면 check_input 이 알려주는지."""
    print("\n[cycle 어긋남 경고]")
    r = run(["scripts/check_input.py", FIX, "--as-of", "2026-09-06"])
    check("check_input 실행", r.returncode in (0, 1))
    check("cycle 검사가 결과를 냄", "cycle" in r.stdout, r.stdout.strip()[-160:])


def main():
    print("=" * 68)
    print("엣지케이스 테스트")
    print("=" * 68)
    for fn in [t_biz, t_months, t_empty_frames, t_kind_detection,
               t_offset_no_miss_match, t_run_variants, t_single_kind,
               t_empty_vendors, t_broken_json, t_missing_vendor_in_data,
               t_state_cycle,
               t_month_boundary, t_strict_current_month, t_until_last_year,
               t_sheet_headers, t_duplicate_uploads, t_validate_strict,
               t_december, t_as_of_clipping, t_validate_catches_regressions,
               t_config_single_source, t_cycle_warning]:
        try:
            fn()
        except Exception as e:
            print(f"  FAIL  {fn.__name__} 예외: {type(e).__name__}: {e}")
            fails.append(fn.__name__)
    print("\n" + "=" * 68)
    print(f"실패 {len(fails)}건" + (f": {fails}" if fails else " — 전부 통과"))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
