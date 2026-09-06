"""
매입 전자세금계산서·전자계산서 수취 점검 리포트 생성.

설계 원칙 세 가지:
  1. 판정 대상은 config/vendors.json 이지 업로드 데이터가 아니다.
     데이터를 groupby 하면 '한 건도 없는 거래처'를 원리적으로 못 잡는다.
  2. 월 범위·기준일은 실행 시점에서 계산한다. 고정값 금지.
  3. 발급기한(익월 10일)이 안 지난 달은 미수취로 확정하지 않는다.

사용법:
    python scripts/build_report.py [--uploads DIR] [--out DIR] [--as-of YYYY-MM-DD]
                                   [--strict] [--no-state]
"""
import argparse
import json
import os
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common import (SKILL_DIR, _load_json, load_config, load_vendors, load_uploads,
                    match_offsets, live_rows, norm_biz, fmt_biz,
                    month_range, month_label, clip_period,
                    is_in_grace, deadline_for, vat_period_of)

GRADE_ORDER = {"확인 필요": 0, "기한 전": 1, "단발·비정기": 2, "정상": 3,
               "거래 종료": 4, "해결": 5}


# ================================================================ 집계

def build_facts(df, vendors, as_of, cfg):
    """
    거래처별 사실관계를 만든다. vendors 를 왼쪽에 두고 데이터를 붙이므로
    데이터에 한 건도 없는 거래처도 반드시 결과에 포함된다.
    """
    months = month_range(as_of)
    mine = cfg["my_biz"]["biz_no"]
    df = clip_period(df[df["받는자번호"] == mine], months, as_of)

    by_vendor = {sid: g for sid, g in df.groupby("공급자번호")}
    listed = {norm_biz(v["biz_no"]) for v in vendors}

    facts = []
    for v in vendors:
        sid = norm_biz(v["biz_no"])
        g = by_vendor.get(sid, df.iloc[0:0])
        facts.append(_one(sid, v, g, months, as_of, cfg, listed=True))

    for sid, g in by_vendor.items():
        if sid in listed:
            continue
        v = {"biz_no": sid, "name": str(g["상호"].iloc[-1]), "cycle": "미등록", "note": ""}
        facts.append(_one(sid, v, g, months, as_of, cfg, listed=False))

    return facts, months


def _one(sid, v, g, months, as_of, cfg, listed):
    cells, got, missing, cancelled, grace = {}, [], [], [], []
    # 연중에 거래를 시작·종료한 거래처는 그 밖의 달을 결번으로 세지 않는다.
    # 점검 기간에 전년 12월이 섞일 수 있으므로 범위도 연도별로 따로 계산한다.
    scope = {y: (_m(v.get("since"), y, 1), _m(v.get("until"), y, 12))
             for y in {yy for yy, _ in months}}

    for ym in months:
        y, m = ym
        gm = g[(g["연"] == y) & (g["월"] == m)] if len(g) else g
        in_grace = is_in_grace(y, m, as_of, cfg)
        lo, hi = scope[y]
        out_of_scope = not (lo <= m <= hi)
        if len(gm) == 0:
            state = "out" if out_of_scope else ("grace" if in_grace else "none")
            cells[ym] = {"amount": None, "state": state}
            if state == "grace":
                grace.append(ym)
            elif state == "none":
                missing.append(ym)
        else:
            lv = live_rows(gm)
            state = "cancelled" if len(lv) == 0 else "ok"
            cells[ym] = {"amount": int(gm["공급가액n"].sum()), "state": state,
                         "count": len(gm)}
            (cancelled if state == "cancelled" else got).append(ym)

    lv_all = live_rows(g)
    if len(lv_all):
        last = lv_all["작성일"].max().date()
        elapsed = (as_of - last).days
    else:
        # 유효 건이 없으면 취소된 날짜를 최종수취일로 쓰지 않는다.
        # 그럴듯한 오답이 크래시보다 나쁘다.
        last, elapsed = None, None

    return {
        "biz_no": sid, "name": v["name"], "cycle": v.get("cycle", ""),
        "note": v.get("note", ""), "listed": listed,
        "cells": cells, "got": got, "missing": missing,
        "cancelled": cancelled, "grace": grace,
        "months_got": len(got), "rows": len(g), "ended": v.get("until") or "",
        "net": int(g["공급가액n"].sum()) if len(g) else 0,
        "last": last, "elapsed": elapsed,
    }


# ================================================================ 판정

def grade_all(facts, as_of, cfg):
    """
    A = 정기 거래처인데 결번 / 발행 후 전액취소 / 한 건도 없음
    C = 최종수취 후 임계일 경과 (A 와 중복되면 A 만 남긴다)
    """
    th = cfg["thresholds"]
    rows, a_ids = [], set()

    for f in facts:
        if not f["listed"]:
            continue
        monthly = f["cycle"] == "매월" or f["months_got"] >= th["monthly_min_months"]

        # 수취 0건이어도 '기한이 지난 결번'이 하나도 없으면 아직 판정할 게 없다.
        # f["missing"] 은 유예 중인 달과 since/until 밖의 달을 이미 빼고 남은 달이다.
        # 이 조건이 없으면 1월 초 실행에서 전 거래처가 '한 건도 없음'으로 잡힌다
        # (1/5 기준 28곳 중 27곳). 2월 초에는 since 가 미래인 거래처가 잡힌다.
        if f["rows"] == 0:
            if not f["missing"]:
                continue
            rows.append(_row("A. 한 건도 없음", "확인 필요", f, f["missing"],
                             f"등록 주기 '{f['cycle']}' 인데 {_labels(f['missing'], as_of)} 수취 0건 · "
                             f"사업자번호 오타 또는 거래 종료 확인"))
            a_ids.add(f["biz_no"])
        elif f["cancelled"]:
            rows.append(_row("A. 발행 후 전액취소", "확인 필요", f, f["cancelled"],
                             "수정세금계산서로 전액 상계 → 재발행 요청 필요"))
            a_ids.add(f["biz_no"])
        elif monthly and f["missing"]:
            rows.append(_row("A. 정기 거래처 결번", "확인 필요", f, f["missing"],
                             f"{_pat(f, as_of)} · 정기 패턴 대비 결번"))
            a_ids.add(f["biz_no"])

    for f in facts:
        if not f["listed"] or f["biz_no"] in a_ids:
            continue
        if f["elapsed"] is None or f["elapsed"] < th["stale_days"]:
            continue

        # 실행 당월은 언제나 유예 상태라 판정에 쓰면 안 된다.
        # 실제로 기다리는 중인 달은 '지난 달인데 아직 기한이 안 지난' 달뿐이다.
        pending = [ym for ym in f["grace"] if ym != (as_of.year, as_of.month)]

        if f["ended"]:
            grade, memo, gap = "거래 종료", \
                f"{_pat(f, as_of)} · vendors.json 에 until={f['ended']} 로 등록됨 — 점검 대상 아님", []
        elif pending and not f["missing"]:
            grade, memo, gap = "기한 전", \
                f"{_pat(f, as_of)} · {_grace_note(pending, as_of, cfg)}", []
        elif not f["missing"] and not f["cancelled"]:
            grade, memo, gap = "정상", f"{_pat(f, as_of)} · 결번 없음 · 발행일이 월초라 당월분 미도래", []
        elif f["months_got"] <= cfg["thresholds"]["sporadic_max_months"]:
            grade, memo, gap = "단발·비정기", f"{_pat(f, as_of)} · 정기 거래 아님 → 경과일 기준 무의미", []
        else:
            grade, memo, gap = "확인 필요", f"{_pat(f, as_of)} · 연속 수취 후 중단됨", f["missing"]
        rows.append(_row(f"C. 최종수취 {cfg['thresholds']['stale_days']}일 경과",
                         grade, f, gap, memo))

    rows.sort(key=lambda r: (GRADE_ORDER.get(r["grade"], 9),
                             r["last"] or date(1900, 1, 1)))
    return rows


def _m(iso, year, default):
    """
    since/until 값(YYYY-MM 또는 YYYY-MM-DD)을 당해년도 월 번호로 환산한다.
    `default=1` 이면 하한(since), `default=12` 면 상한(until)으로 쓰인다.

    다른 해일 때 경계값을 잘못 주면 조용히 틀린다. 특히 **작년에 끝난 거래처**
    (until=2025-08) 에 상한 12 를 주면 올해 전 월이 점검 범위가 되어
    '한 건도 없음 · 확인 필요' 로 매번 잡힌다 — until 을 넣으라는 안내와 정반대다.
    """
    if not iso:
        return default
    y, m = int(str(iso)[:4]), int(str(iso)[5:7])
    if y == year:
        return m
    is_lower = (default == 1)
    if y < year:
        # 올해 이전: 하한이면 1월부터 포함, 상한이면 올해는 통째로 범위 밖(0)
        return 1 if is_lower else 0
    # 올해 이후: 하한이면 올해는 통째로 범위 밖(13), 상한이면 12월까지 포함
    return 13 if is_lower else 12


def _labels(yms, as_of):
    """(연,월) 목록을 '1월, 2월' / '전년 12월' 형태로."""
    return ", ".join(month_label(ym, as_of) for ym in yms)


def _row(kind, grade, f, gap, memo):
    return {"kind": kind, "grade": grade, "biz_no": f["biz_no"], "name": f["name"],
            "cycle": f["cycle"], "gap": gap, "last": f["last"], "memo": memo,
            "note": f["note"]}


def _pat(f, as_of):
    if not f["got"]:
        return "수취 0개월"
    return f"{len(f['got'])}개월 수취({_labels(f['got'], as_of)})"


def _grace_note(pending, as_of, cfg):
    ym = min(pending)
    dl = deadline_for(ym[0], ym[1], cfg["grace"]["deadline_day"])
    left = (dl - as_of).days
    return f"{month_label(ym, as_of)}분 발급기한 {dl} (D-{left}) — 아직 미수취 아님"


# ================================================================ 상태 비교

def load_state(cfg):
    p = os.path.join(SKILL_DIR, cfg["state"]["path"])
    if not os.path.exists(p):
        return None
    return _load_json(p, "실행 이력(last-run.json)",
                      "이 파일만 지우고 다시 실행하면 이번 회차부터 다시 쌓인다. "
                      "다만 '신규/계속/해결' 비교는 한 번 끊긴다.")


def apply_state(rows, prev, as_of, cfg):
    """
    지난 실행과 비교해 신규 / 계속 / 해결 을 붙인다.
    월 1~2회 주기에서는 '지난달 요청한 게 해결됐나'가 가장 알고 싶은 정보다.
    """
    if not cfg["state"]["enabled"]:
        for r in rows:
            r["status"] = ""
        return rows

    prev_map = {p["biz_no"]: p for p in (prev or {}).get("flagged", [])}
    now_ids = {r["biz_no"] for r in rows if r["grade"] == "확인 필요"}

    for r in rows:
        p = prev_map.get(r["biz_no"])
        if r["grade"] != "확인 필요":
            r["status"] = ""
            r["first_seen"] = None
        elif p:
            r["first_seen"] = p.get("first_seen")
            n = _months_between(p.get("first_seen"), as_of)
            r["status"] = f"계속 ({n}개월)" if n else "계속"
        else:
            r["first_seen"] = as_of.isoformat()
            r["status"] = "신규"

    if cfg["state"]["resolved_shown_once"]:
        for bid, p in prev_map.items():
            if bid not in now_ids:
                rows.append({"kind": "해결됨", "grade": "해결", "biz_no": bid,
                             "name": p.get("name", ""), "cycle": "", "gap": [],
                             "last": None, "status": "해결",
                             "memo": f"직전 실행({(prev or {}).get('run_date','')})에서 '{p.get('reason','')}' 로 잡혔으나 이번엔 정상",
                             "note": "", "first_seen": None})
    return rows


def _months_between(iso, as_of):
    if not iso:
        return 0
    d = date.fromisoformat(iso)
    return (as_of.year - d.year) * 12 + (as_of.month - d.month)


def save_state(rows, as_of, cfg):
    if not cfg["state"]["enabled"]:
        return
    flagged = [{"biz_no": r["biz_no"], "name": r["name"], "grade": r["grade"],
                "reason": r["kind"], "first_seen": r.get("first_seen") or as_of.isoformat()}
               for r in rows if r["grade"] == "확인 필요"]
    p = os.path.join(SKILL_DIR, cfg["state"]["path"])
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        json.dump({"run_date": as_of.isoformat(), "flagged": flagged},
                  f, ensure_ascii=False, indent=2)


# ================================================================ 엑셀

class Style:
    def __init__(self, cfg):
        c, fm = cfg["colors"], cfg["format"]
        fill = lambda k: PatternFill("solid", fgColor=c[k]) if c.get(k) else None
        self.hdr = PatternFill("solid", fgColor=c["header_bg"])
        self.hdr_font = Font(name=fm["font"], color=c["header_font"], bold=True, size=10)
        self.body = Font(name=fm["font"], size=10)
        self.title = Font(name=fm["font"], bold=True, size=13)
        self.sub = Font(name=fm["font"], size=9, italic=True)
        self.border = Border(*[Side("thin", color=c["border"])] * 4)
        self.cancelled = fill("cancelled")
        self.money, self.date = fm["money"], fm["date"]
        self.by_grade = {
            "확인 필요": fill("grade_action"), "기한 전": fill("grade_grace"),
            "정상": fill("grade_normal"), "단발·비정기": fill("grade_sporadic"),
            "해결": fill("grade_normal"), "거래 종료": fill("grade_normal"),
        }

    def header(self, ws, row, n):
        for i in range(1, n + 1):
            c = ws.cell(row, i)
            c.fill, c.font = self.hdr, self.hdr_font
            c.alignment = Alignment("center", "center", wrap_text=True)


def write_matrix(wb, st, facts, months, as_of, cfg):
    ws = wb.active
    ws.title = cfg["sheets"]["matrix"]
    mb = cfg["my_biz"]

    ws["A1"] = f"공급자 × 월 매입 발급 매트릭스 ({as_of.year}-01-01 ~ {as_of})"
    ws["A1"].font = st.title
    ws["A2"] = (f"공급받는자: {fmt_biz(mb['biz_no'])} {mb['name']} · 값=월별 순공급가액(원) · "
                f"빈칸=미수취 · 노랑=발행 후 전액취소 · 연파랑=발급기한 전")
    ws["A2"].font = st.sub

    head = ["공급자번호", "상호", "주기"] + [month_label(ym, as_of) for ym in months] + \
           ["순공급가액", "최종수취일", "대상"]
    ws.append([])
    if cfg["vat_period"]["enabled"]:
        band = ["", "", ""] + [vat_period_of(m, cfg) for _, m in months] + ["", "", ""]
        ws.append(band)
        for i in range(4, 4 + len(months)):
            ws.cell(4, i).font = st.sub
            ws.cell(4, i).alignment = Alignment("center")
        hdr_row = 5
    else:
        hdr_row = 4
    ws.append(head)
    st.header(ws, hdr_row, len(head))

    facts = sorted(facts, key=lambda f: (not f["listed"], -f["months_got"], f["name"]))
    for f in facts:
        line = [fmt_biz(f["biz_no"]), f["name"], f["cycle"]]
        for ym in months:
            line.append(f["cells"][ym]["amount"])
        line += [f["net"], f["last"], "점검대상" if f["listed"] else "대상외"]
        ws.append(line)
        r = ws.max_row
        for i, ym in enumerate(months):
            s = f["cells"][ym]["state"]
            if s == "cancelled":
                ws.cell(r, 4 + i).fill = st.cancelled
            elif s == "grace" and st.by_grade["기한 전"]:
                ws.cell(r, 4 + i).fill = st.by_grade["기한 전"]

    for row in ws.iter_rows(min_row=hdr_row + 1, max_row=ws.max_row, max_col=len(head)):
        for c in row:
            c.font, c.border = st.body, st.border
            c.alignment = Alignment(vertical="center", horizontal="right")
        row[0].alignment = Alignment("center", "center")
        row[1].alignment = Alignment("left", "center")
        row[2].alignment = Alignment("center", "center")
        for i in range(3, 3 + len(months) + 1):
            row[i].number_format = st.money
        row[3 + len(months) + 1].number_format = st.date
        row[-1].alignment = Alignment("center", "center")

    ws.freeze_panes = f"D{hdr_row + 1}"
    widths = [14, 28, 9] + [13] * len(months) + [14, 13, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def write_missing(wb, st, rows, cfg, dedup_count, as_of):
    ws = wb.create_sheet(cfg["sheets"]["missing"])
    use_state = cfg["state"]["enabled"]
    head = ["유형", "등급"] + (["상태"] if use_state else []) + \
           ["공급자번호", "상호", "주기", "누락된 월", "최종수취일", "메모", "비고"]

    n = {g: sum(1 for r in rows if r["grade"] == g) for g in GRADE_ORDER}
    ws["A1"] = (f"미수취 의심 목록 · 확인 필요 {n.get('확인 필요',0)}건 / "
                f"기한 전 {n.get('기한 전',0)}건 / 단발·비정기 {n.get('단발·비정기',0)}건 / "
                f"정상 {n.get('정상',0)}건 / 거래 종료 {n.get('거래 종료',0)}건 "
                f"(A·C 중복 {dedup_count}곳 제외)")
    ws["A1"].font = st.title
    ws["A2"] = ("등급: 확인 필요 → 실제 조치 대상 / 기한 전 → 발급기한 남음, 대기 / "
                "단발·비정기 → 정기거래 아님 / 정상 → 결번 없음, 무시 가능")
    ws["A2"].font = st.sub
    ws.append([])
    # 헤더 행 번호를 고정값으로 적으면 제목 줄이 한 줄만 늘어도 어긋난다.
    # (실제로 서식은 3행에, 헤더는 4행에 앉아 있었다.)
    ws.append(head)
    hdr_row = ws.max_row          # append 직후에 읽어야 정확하다
    st.header(ws, hdr_row, len(head))

    for r in rows:
        line = [r["kind"], r["grade"]] + ([r.get("status", "")] if use_state else []) + \
               [fmt_biz(r["biz_no"]), r["name"], r["cycle"],
                _labels(r["gap"], as_of), r["last"], r["memo"], r["note"]]
        ws.append(line)
        fill = st.by_grade.get(r["grade"])
        if fill:
            for c in ws[ws.max_row]:
                c.fill = fill

    di = 7 if use_state else 6
    for row in ws.iter_rows(min_row=hdr_row + 1, max_row=ws.max_row, max_col=len(head)):
        for c in row:
            c.font, c.border = st.body, st.border
            c.alignment = Alignment(vertical="center", wrap_text=True)
        row[di].number_format = st.date

    widths = [20, 11] + ([12] if use_state else []) + [14, 26, 9, 26, 13, 50, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    # 헤더 '아래'를 얼려야 헤더가 남는다. 헤더 행 자체를 얼면 스크롤할 때 사라진다.
    ws.freeze_panes = f"A{hdr_row + 1}"


def unlisted_reco(f, cfg):
    """
    vendors.json 에 추가할지 권할 기준. 임계값은 config 가 유일한 출처다.

    월수만 보면 '한 달에 몰아서 큰 금액'이 통째로 묻힌다(7·8월에만 4,953만원이
    찍힌 곳이 아무 표시 없이 지나갔다). 그래서 금액 기준을 함께 본다.
    """
    th = cfg["thresholds"]
    if f["months_got"] >= th["unlisted_recommend_months"]:
        return "추가 권장 (정기성)"
    if f["net"] >= th["unlisted_recommend_amount"]:
        return "추가 권장 (금액)"
    return ""


def write_unlisted(wb, st, facts, cfg, as_of):
    """점검 대상에는 없는데 자료에 등장한 공급자."""
    ws = wb.create_sheet(cfg["sheets"]["unlisted"])
    un = [f for f in facts if not f["listed"]]
    ws["A1"] = f"점검 대상에 없는데 자료에 등장한 공급자 · {len(un)}곳"
    ws["A1"].font = st.title
    ws["A2"] = "정기 거래가 된 곳은 config/vendors.json 에 추가할 것. 일회성이면 그대로 둔다."
    ws["A2"].font = st.sub
    ws.append([])
    head = ["공급자번호", "상호", "수취 월", "건수", "순공급가액", "최종수취일", "추가 권장"]
    ws.append(head)
    hdr_row = ws.max_row          # append 직후에 읽어야 정확하다
    st.header(ws, hdr_row, len(head))

    for f in sorted(un, key=lambda x: (-x["months_got"], -x["net"])):
        ws.append([fmt_biz(f["biz_no"]), f["name"],
                   _labels(f["got"], as_of) if f["got"] else "",
                   f["rows"], f["net"], f["last"], unlisted_reco(f, cfg)])
    for row in ws.iter_rows(min_row=hdr_row + 1, max_row=ws.max_row, max_col=len(head)):
        for c in row:
            c.font, c.border = st.body, st.border
        row[4].number_format = st.money
        row[5].number_format = st.date
    for i, w in enumerate([14, 30, 22, 8, 14, 13, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{hdr_row + 1}"


def write_raw(wb, st, df, cfg):
    ws = wb.create_sheet(cfg["sheets"]["raw"])
    head = ["작성일자", "구분", "승인번호", "공급자번호", "공급자상호", "공급받는자번호",
            "공급받는자상호", "품목명", "합계금액", "공급가액", "종류", "발급유형",
            "상계처리", "원본파일"]
    ws.append(head)
    hdr_row = ws.max_row          # append 직후에 읽어야 정확하다
    st.header(ws, hdr_row, len(head))
    d = df.copy()
    d["상계"] = d["상계"].replace("", "-")
    for _, r in d.iterrows():
        ws.append([r["작성일자"], r["구분"], r["승인번호"], fmt_biz(r["공급자번호"]),
                   r["상호"], fmt_biz(r["받는자번호"]), r.get("공급받는자상호", ""),
                   r.get("품목명", ""), int(r["합계금액n"]), int(r["공급가액n"]),
                   r.get("전자세금계산서종류", ""), r.get("발급유형", ""),
                   r["상계"], r["원본파일"]])
        if r["상계"] != "" and st.cancelled and "상계" in str(r["상계"]):
            for c in ws[ws.max_row]:
                c.fill = st.cancelled
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(head)):
        for c in row:
            c.font = st.body
        row[8].number_format = st.money
        row[9].number_format = st.money
    for i, w in enumerate([12, 6, 30, 14, 28, 14, 28, 30, 13, 13, 11, 11, 20, 32], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{hdr_row + 1}"


# ================================================================ main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--uploads", default="/mnt/user-data/uploads")
    ap.add_argument("--out", default="/mnt/user-data/outputs")
    ap.add_argument("--as-of", default=None, help="기준일. 미지정시 오늘.")
    ap.add_argument("--strict", action="store_true",
                    help="발급기한 유예 없이 전부 미수취로 판정 (부가세 신고 직전용)")
    ap.add_argument("--no-state", action="store_true", help="직전 실행 비교 끄기")
    a = ap.parse_args()

    as_of = date.fromisoformat(a.as_of) if a.as_of else date.today()
    cfg = load_config()
    if a.strict:
        cfg["grace"]["mode"] = "strict"
    if a.no_state:
        cfg["state"]["enabled"] = False

    vendors = load_vendors()
    df_all, _meta = load_uploads(a.uploads, cfg)
    mine = cfg["my_biz"]["biz_no"]
    excluded = int((df_all["받는자번호"] != mine).sum())
    pairs = match_offsets(df_all, cfg)

    # 리포트의 모든 시트가 같은 기간을 본다. 원본정제데이터만 안 자르면
    # 검증기가 다른 기준으로 세어 멀쩡한 산출물에 FAIL 이 난다.
    months = month_range(as_of)
    mine_rows = df_all[df_all["받는자번호"] == mine]
    df = clip_period(mine_rows, months, as_of)
    out_of_period = len(mine_rows) - len(df)

    facts, months = build_facts(df, vendors, as_of, cfg)
    rows = grade_all(facts, as_of, cfg)
    a_ids = {r["biz_no"] for r in rows if r["kind"].startswith("A")}
    dedup = sum(1 for f in facts
                if f["listed"] and f["biz_no"] in a_ids
                and f["elapsed"] is not None
                and f["elapsed"] >= cfg["thresholds"]["stale_days"])

    prev = load_state(cfg)
    rows = apply_state(rows, prev, as_of, cfg)

    st = Style(cfg)
    wb = Workbook()
    write_matrix(wb, st, facts, months, as_of, cfg)
    write_missing(wb, st, rows, cfg, dedup, as_of)
    write_unlisted(wb, st, facts, cfg, as_of)
    write_raw(wb, st, df, cfg)

    os.makedirs(a.out, exist_ok=True)
    name = cfg["output"]["filename_template"].format(
        year=as_of.year, asof=as_of.strftime("%m%d"))
    path = os.path.join(a.out, name)
    wb.save(path)
    save_state(rows, as_of, cfg)

    n = {g: sum(1 for r in rows if r["grade"] == g) for g in list(GRADE_ORDER) + ["해결"]}
    span = f"{month_label(months[0], as_of)}~{month_label(months[-1], as_of)}"
    print(f"기준일 {as_of} · 대상 월 {span} · 모드 {cfg['grace']['mode']}")
    print(f"총 {len(df)}건 / 매출분 제외 {excluded}건 / 기간 밖 제외 {out_of_period}건 / "
          f"상계쌍 {pairs}쌍")
    print(f"점검대상 {sum(1 for f in facts if f['listed'])}곳 / "
          f"대상외 {sum(1 for f in facts if not f['listed'])}곳")
    print(f"확인 필요 {n.get('확인 필요',0)} · 기한 전 {n.get('기한 전',0)} · "
          f"단발 {n.get('단발·비정기',0)} · 정상 {n.get('정상',0)} · 해결 {n.get('해결',0)}")
    _print_deadline(facts, as_of, cfg)
    _print_reco(facts, cfg)
    print(f"저장: {path}")
    return path


def _print_deadline(facts, as_of, cfg):
    """
    발급기한 D-day. 지금까지는 메모 문자열 안에만 있어서, 리포트를 열지 않으면
    '며칠 남았는지'가 전혀 안 보였다. 월 1~2회 실행에서 가장 먼저 봐야 할 숫자다.
    """
    pend = sorted({ym for f in facts if f["listed"] for ym in f["grace"]
                   if ym != (as_of.year, as_of.month)})
    if not pend:
        return
    ym = pend[0]
    dl = deadline_for(ym[0], ym[1], cfg["grace"]["deadline_day"])
    waiting = sum(1 for f in facts if f["listed"] and ym in f["grace"])
    print(f"발급기한 임박: {month_label(ym, as_of)}분 기한 {dl} "
          f"(D-{(dl - as_of).days}) · 아직 미수취 {waiting}곳")


def _print_reco(facts, cfg):
    """대상외 중 '추가 권장'. 시트를 안 열면 놓치므로 콘솔에도 띄운다."""
    un = [f for f in facts if not f["listed"] and unlisted_reco(f, cfg)]
    if not un:
        return
    print(f"vendors.json 추가 권장 {len(un)}곳:")
    for f in sorted(un, key=lambda x: -x["net"]):
        print(f"  {fmt_biz(f['biz_no'])} {f['name']} · {f['net']:,}원 · "
              f"{unlisted_reco(f, cfg)}")


if __name__ == "__main__":
    main()
