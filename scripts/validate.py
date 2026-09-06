"""
산출물 xlsx 자체 검증. build_report.py 직후에 반드시 돌린다.

설계 원칙: **0건 매칭이면 PASS 가 아니라 FAIL 이다.**
"미수취 칸에 색이 칠해져 있지 않다"는, 정말 없어서일 수도 있지만
셀을 못 찾아서일 수도 있다. 후자를 PASS 로 넘기면 검증이 무의미해진다.
그래서 모든 검사는 '검사 대상을 몇 개 찾았는지'를 먼저 보고한다.

열은 반드시 헤더 '이름'으로 찾는다. 인덱스 하드코딩 금지 —
열이 하나만 추가돼도 엉뚱한 칸을 검사하게 된다.

사용법:
    python scripts/validate.py <출력.xlsx> [--uploads DIR] [--as-of YYYY-MM-DD] [--strict]

build_report.py 에 준 `--as-of` / `--strict` 는 **여기에도 똑같이 준다.**
안 주면 검증기가 다른 기준으로 판단해서 멀쩡한 산출물에 FAIL 을 낸다.
"""
import argparse
import os
import sys
from datetime import date

import pandas as pd
from openpyxl import load_workbook

from common import (SKILL_DIR, load_config, load_vendors, load_uploads,
                    match_offsets, norm_biz, month_range, is_in_grace)

PASS, FAIL, SKIP = "PASS", "FAIL", "SKIP"
_res = []


def rec(level, name, detail=""):
    _res.append((level, name, detail))


def hdr_map(ws, row):
    """헤더 행에서 {이름: 열번호}. 위치 대신 이름으로 찾기 위한 것."""
    return {str(c.value).strip(): c.column for c in ws[row] if c.value}


def find_header_row(ws, must_have, limit=8):
    for r in range(1, min(limit, ws.max_row) + 1):
        m = hdr_map(ws, r)
        if all(k in m for k in must_have):
            return r, m
    return None, {}


# ---------------------------------------------------------------- 검사

def check_sheets(wb, cfg):
    # config 의 _comment 키는 설정값이 아니다.
    want = [v for k, v in cfg["sheets"].items() if not k.startswith("_")]
    missing = [s for s in want if s not in wb.sheetnames]
    if missing:
        rec(FAIL, "시트 구성", f"없는 시트: {missing} / 실제: {wb.sheetnames}")
    else:
        rec(PASS, "시트 구성", " · ".join(wb.sheetnames))


def check_matrix_months(wb, cfg, as_of):
    ws = wb[cfg["sheets"]["matrix"]]
    want = [f"{m}월" for m in month_range(as_of)]
    r, m = find_header_row(ws, ["공급자번호", "상호"])
    if not r:
        rec(FAIL, "매트릭스 헤더", "'공급자번호'/'상호' 헤더 행을 못 찾음")
        return None, None
    found = [k for k in m if k.endswith("월")]
    if found != want:
        rec(FAIL, "매트릭스 월 컬럼",
            f"기대 {want}\n실제 {found}\n월 범위가 고정값으로 박혀 있을 수 있다.")
    else:
        rec(PASS, "매트릭스 월 컬럼", f"1~{as_of.month}월 · {len(want)}개")
    return r, m


def check_matrix_blank_and_fill(wb, cfg, as_of):
    """
    미수취 = 값 없음 + 채우기 없음.
    전액취소 = 값 0 + 노랑.
    검사 대상을 하나도 못 찾으면 FAIL (마크업이 바뀐 것).
    """
    ws = wb[cfg["sheets"]["matrix"]]
    r, m = find_header_row(ws, ["공급자번호", "상호"])
    if not r:
        return
    mcols = [m[f"{x}월"] for x in month_range(as_of) if f"{x}월" in m]
    if not mcols:
        rec(FAIL, "매트릭스 셀 검사", "월 컬럼을 하나도 못 찾음 — 이후 검사 무의미")
        return

    blank, cancelled, bad_blank, bad_zero = 0, 0, [], []
    yellow = (cfg["colors"]["cancelled"] or "").upper()
    for row in ws.iter_rows(min_row=r + 1, max_row=ws.max_row):
        for ci in mcols:
            c = ws.cell(row[0].row, ci)
            fill = (c.fill.fgColor.rgb or "") if c.fill and c.fill.fill_type else ""
            fill = str(fill).upper()[-6:]
            if c.value is None:
                blank += 1
                if fill == yellow:
                    bad_blank.append(c.coordinate)
            elif c.value == 0:
                cancelled += 1
                if fill != yellow:
                    bad_zero.append(c.coordinate)

    if blank == 0:
        rec(FAIL, "미수취=공란 검사", "공란 셀을 하나도 못 찾음. 셀 값 표기 방식이 바뀌었을 수 있다.")
    elif bad_blank:
        rec(FAIL, "미수취=공란 검사", f"공란인데 취소색이 칠해진 셀: {bad_blank[:10]}")
    else:
        rec(PASS, "미수취=공란 검사", f"공란 {blank}칸 확인 · 잘못된 채우기 없음")

    if cancelled == 0:
        rec(SKIP, "전액취소=0+노랑 검사", "값 0 인 셀이 없음 (이번 데이터에 전액취소 없음)")
    elif bad_zero:
        rec(FAIL, "전액취소=0+노랑 검사", f"0 인데 노랑이 아닌 셀: {bad_zero[:10]}")
    else:
        rec(PASS, "전액취소=0+노랑 검사", f"{cancelled}칸 전부 노랑")


def check_vendor_coverage(wb, cfg, vendors):
    """vendors.json 의 모든 거래처가 매트릭스에 행으로 존재해야 한다."""
    ws = wb[cfg["sheets"]["matrix"]]
    r, m = find_header_row(ws, ["공급자번호", "상호"])
    if not r:
        return
    col = m["공급자번호"]
    seen = {norm_biz(ws.cell(i, col).value)
            for i in range(r + 1, ws.max_row + 1) if ws.cell(i, col).value}
    want = {norm_biz(v["biz_no"]) for v in vendors}
    miss = want - seen
    if not vendors:
        rec(SKIP, "점검 대상 누락 검사", "vendors.json 이 비어 있음")
    elif miss:
        rec(FAIL, "점검 대상 누락 검사",
            f"매트릭스에 행이 없는 거래처 {len(miss)}곳: {sorted(miss)[:10]}\n"
            "데이터를 groupby 해서 만들면 '0건 거래처'가 사라진다.")
    else:
        rec(PASS, "점검 대상 누락 검사", f"{len(want)}곳 전부 행 존재")


def check_totals(wb, cfg, df):
    """매트릭스 월별 합 == 원본정제데이터 합. 집계 버그를 잡는다."""
    ws = wb[cfg["sheets"]["matrix"]]
    r, m = find_header_row(ws, ["공급자번호", "상호"])
    if not r or "순공급가액" not in m:
        rec(FAIL, "합계 대사", "'순공급가액' 열을 못 찾음")
        return
    col = m["순공급가액"]
    total = sum(ws.cell(i, col).value or 0 for i in range(r + 1, ws.max_row + 1))
    raw = int(df["공급가액n"].sum())
    if total != raw:
        rec(FAIL, "합계 대사", f"매트릭스 {total:,} vs 원본 {raw:,} · 차 {total-raw:,}")
    else:
        rec(PASS, "합계 대사", f"{total:,}원 일치")


def check_raw_rows(wb, cfg, df):
    ws = wb[cfg["sheets"]["raw"]]
    got = ws.max_row - 1
    if got != len(df):
        rec(FAIL, "원본 행수", f"시트 {got}행 vs 데이터 {len(df)}건")
    else:
        rec(PASS, "원본 행수", f"{got}건")


def check_no_duplicates(wb, cfg):
    """
    원본정제데이터의 승인번호 중복. 같은 구간 파일을 두 번 올린 경우를 잡는다.

    check_input.py 에도 같은 검사가 있지만, 그 단계를 건너뛰면 리포트가
    오류 없이 만들어지고 금액이 부풀려진 채 그대로 나간다(441건 → 521건,
    상계쌍 6 → 8쌍). 산출물만 보고도 잡히도록 여기서 한 번 더 본다.
    원본 행수 검사는 산출물을 자기 입력과 비교하므로 이 경우를 못 잡는다.
    """
    ws = wb[cfg["sheets"]["raw"]]
    r, m = find_header_row(ws, ["승인번호"])
    if not r:
        rec(FAIL, "승인번호 중복", "'승인번호' 열을 못 찾음 — 중복 검사를 못 했다")
        return
    col = m["승인번호"]
    seen, dup = set(), []
    for i in range(r + 1, ws.max_row + 1):
        v = ws.cell(i, col).value
        if v is None or str(v).strip() == "":
            continue
        v = str(v).strip()
        if v in seen:
            dup.append(v)
        seen.add(v)
    if not seen:
        rec(FAIL, "승인번호 중복", "승인번호가 한 건도 없음 — 검사 대상을 못 찾았다")
    elif dup:
        rec(FAIL, f"승인번호 중복 {len(dup)}건",
            "같은 구간 홈택스 파일을 두 번 올렸을 가능성이 높다.\n"
            "그대로 두면 금액과 상계쌍이 부풀려진다. 중복분을 빼고 다시 만들 것.\n" +
            "\n".join(dup[:10]))
    else:
        rec(PASS, "승인번호 중복", f"{len(seen)}건 전부 고유")


def check_grace(wb, cfg, as_of):
    """유예 판정이 실제로 작동했는지. auto 모드인데 유예 칸이 0이면 의심."""
    if cfg["grace"]["mode"] != "auto":
        rec(SKIP, "발급기한 유예", "strict 모드")
        return
    ws = wb[cfg["sheets"]["missing"]]
    r, m = find_header_row(ws, ["등급"])
    if not r:
        rec(FAIL, "발급기한 유예", "'등급' 열을 못 찾음")
        return
    col = m["등급"]
    n = sum(1 for i in range(r + 1, ws.max_row + 1)
            if str(ws.cell(i, col).value) == "기한 전")
    prev = as_of.month - 1
    expect = prev >= 1 and is_in_grace(as_of.year, prev, as_of, cfg)
    if expect and n == 0:
        rec(FAIL, "발급기한 유예",
            f"{prev}월분 기한({cfg['grace']['deadline_day']}일)이 안 지났는데 '기한 전' 판정이 0건. "
            "유예 로직이 안 걸렸을 수 있다.")
    else:
        rec(PASS, "발급기한 유예", f"'기한 전' {n}건 (전월 유예 구간: {expect})")


def check_paths(cfg):
    """
    저장소에서 받아온 파일이 실제로 다 있는지.
    SKILL.md 는 Claude 스킬 폴더에 있고 작업 디렉토리에는 없으므로 제외한다.
    """
    want = ["sync.py", "config/check-config.json", "config/vendors.json",
            "references/column-mapping.md", "references/judgment-rules.md",
            "references/excel-format.md", "scripts/common.py",
            "scripts/check_input.py", "scripts/build_report.py",
            "scripts/validate.py", "scripts/push.py"]
    miss = [p for p in want if not os.path.exists(os.path.join(SKILL_DIR, p))]
    if miss:
        rec(FAIL, "경로 존재 확인",
            f"없는 파일: {miss}\nsync.py 를 다시 돌려 저장소에서 받으세요.")
    else:
        rec(PASS, "경로 존재 확인", f"{len(want)}개 전부 존재")


def check_config_single_source(cfg):
    """
    임계값이 스크립트에 하드코딩돼 있지 않은지 소스에서 직접 찾는다.
    config 가 거짓말을 하기 시작하는 가장 흔한 경로다.
    """
    import re
    th = cfg["thresholds"]
    bad = []
    for fn in ["build_report.py", "check_input.py"]:
        src = open(os.path.join(SKILL_DIR, "scripts", fn), encoding="utf-8").read()
        src = re.sub(r"#.*", "", src)
        for key, val in th.items():
            for pat in [f"range(1, {val})", f"[:{val}]", f">= {val}", f"<= {val}"]:
                if pat in src:
                    bad.append(f"{fn}: '{pat}' — {key}={val} 하드코딩 의심")
    if bad:
        rec(FAIL, "임계값 단일 출처", "\n".join(bad))
    else:
        rec(PASS, "임계값 단일 출처", "스크립트에 임계값 하드코딩 없음")


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--uploads", default="/mnt/user-data/uploads")
    ap.add_argument("--as-of", default=None)
    ap.add_argument("--strict", action="store_true",
                    help="build_report.py 를 --strict 로 만든 산출물을 검증할 때. "
                         "빼먹으면 '기한 전' 0건을 유예 로직 고장으로 오판해 FAIL 이 난다")
    a = ap.parse_args()
    as_of = date.fromisoformat(a.as_of) if a.as_of else date.today()

    cfg = load_config()
    if a.strict:
        cfg["grace"]["mode"] = "strict"
    vendors = load_vendors()
    wb = load_workbook(a.xlsx)
    df, _ = load_uploads(a.uploads, cfg)
    match_offsets(df, cfg)
    df = df[(df["받는자번호"] == cfg["my_biz"]["biz_no"]) & (df["연"] == as_of.year)]

    check_paths(cfg)
    check_config_single_source(cfg)
    check_sheets(wb, cfg)
    check_matrix_months(wb, cfg, as_of)
    check_matrix_blank_and_fill(wb, cfg, as_of)
    check_vendor_coverage(wb, cfg, vendors)
    check_totals(wb, cfg, df)
    check_raw_rows(wb, cfg, df)
    check_no_duplicates(wb, cfg)
    check_grace(wb, cfg, as_of)

    print(f"\n산출물 검증 · {os.path.basename(a.xlsx)} · 기준일 {as_of}\n" + "=" * 68)
    for lv, n, d in _res:
        print(f"[{lv}] {n}")
        if d:
            for line in str(d).splitlines():
                print(f"       {line}")
    f = sum(1 for lv, _, _ in _res if lv == FAIL)
    s = sum(1 for lv, _, _ in _res if lv == SKIP)
    print("=" * 68)
    print(f"FAIL {f} · SKIP {s} · PASS {len(_res)-f-s}")
    if f:
        print("\nFAIL 이 있으면 산출물을 사용자에게 주지 말 것.")
    return 1 if f else 0


if __name__ == "__main__":
    sys.exit(main())
